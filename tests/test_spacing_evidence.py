"""Spacing and local Word styling cannot decide authorship or hide phrase matches."""
from types import SimpleNamespace as NS
import io

import pytest
from docx import Document
from docx.shared import Pt

import backend_adapter as b
import nebras_report as nr
import tech_signals as tech
from matching_text import matching_text_map
from tests.test_detection_layers import _docx, LLM_PARA, HUMAN_PARA, HUMAN_SLIPPY
from tests.test_targeted_highlights import _report, _finding


def test_collapsed_phrase_maps_to_the_complete_original_phrase():
    text = "Prefix: not\u2002\u2002only  clear\tbut also useful.\nNext line."
    normalized, positions = matching_text_map(text)
    a, z = normalized.index("not"), normalized.index("useful") + len("useful")
    assert text[positions[a]:positions[z - 1] + 1] == "not\u2002\u2002only  clear\tbut also useful"
    assert normalized.endswith("\nNext line.")


@pytest.mark.parametrize("spacing", ["  ", "     ", "\u2002", "\t"])
def test_word_spacing_does_not_hide_generated_wording_or_change_source(spacing):
    original = LLM_PARA.replace(" ", spacing)
    plain = b.analyze_uploaded_bytes("plain.docx", _docx([LLM_PARA] * 20))
    spaced = b.analyze_uploaded_bytes("spaced.docx", _docx([original] * 20))
    assert original in spaced.texts  # analysis never rewrites the uploaded source
    assert spaced.features["contrast"][2] == pytest.approx(plain.features["contrast"][2])
    assert spaced.score == pytest.approx(plain.score, abs=1.0)
    assert spaced.score >= 65
    assert spaced.features["err_density"][2] == 0


def test_real_typing_errors_are_local_and_not_erased():
    r = b.analyze_uploaded_bytes("local.docx", _docx([LLM_PARA] * 12 + [HUMAN_SLIPPY]))
    assert r.text_source_error_counts["repeated_word"] > 0
    assert r.features["err_density"][2] < 0
    assert r.text_error_shifts[12] < 0
    assert r.text_error_shifts[0] == r.file_shifts[0] == 0
    assert r.file_shifts[12] < 0


def test_surface_quality_findings_remain_visible_without_human_credit():
    r = b.analyze_uploaded_bytes("spacing.docx", _docx([HUMAN_PARA.replace(" ", "  ")] * 12))
    assert any(f.kind == "double_space" for f in r.findings)
    assert r.features["err_density"][2] == 0
    assert r.features["clean"][2] == 0
    assert r.score < 52


def test_word_table_errors_do_not_give_human_credit_to_the_main_text():
    doc = Document()
    for _ in range(12):
        doc.add_paragraph(LLM_PARA)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = HUMAN_SLIPPY
    buffer = io.BytesIO(); doc.save(buffer)
    r = b.analyze_uploaded_bytes("table.docx", buffer.getvalue())
    assert any(p.role == "table" and HUMAN_SLIPPY in p.text for p in r.doc_paras)
    assert not r.text_source_error_counts
    assert r.features["err_density"][2] == 0


def test_spacing_does_not_suppress_a_complete_supported_sentence():
    first = "This  framework provides evidence for adaptive interventions across institutions and communities through a coordinated approach to implementation."
    text = first + " We measured the water level at eighteen stations on Tuesday and checked each reading against the original notebook before replacing the pump battery."
    trace = NS(para=0, kind="double_space", start=4, end=6, note="")
    r = _report(text, [_finding(text, "This  framework"), trace])
    p = nr.plan_highlights(r, [NS(text=text, role="body")])
    assert p.flags[0] == [(0, len(first))]


def test_minor_word_font_change_does_not_penalize_the_surrounding_paragraph():
    doc = Document()
    for i in range(12):
        p = doc.add_paragraph()
        first = p.add_run(LLM_PARA[:8]); first.font.size = Pt(8 if i == 5 else 12)
        rest = p.add_run(LLM_PARA[8:]); rest.font.size = Pt(12)
    buffer = io.BytesIO(); doc.save(buffer)
    r = b.analyze_uploaded_bytes("font.docx", buffer.getvalue())
    assert any(f.kind == "fmt_size" and f.para == 5 for f in r.findings)
    assert r.file_shifts[5] == r.file_shifts[0]


def test_pooled_word_formatting_credit_uses_the_same_locations_as_the_map():
    from tests.test_structure_fairness import _docx as formatted_docx
    rows = [("Literature Review", "center", 28, True)]
    rows += [(LLM_PARA, "center" if i == 5 else "both", 24) for i in range(12)]
    r = b.analyze_uploaded_bytes("format.docx", formatted_docx(rows))
    value, weight, contribution = r.features["fmt_incons"]
    assert contribution == pytest.approx(value * weight / 12)


def test_technical_phrase_matching_ignores_spacing_but_retains_spelling_errors():
    text = "The robust controller runs within the evaluated framework, providing the result to the next stage."
    plain = tech.measure([text] * 80)
    spaced = tech.measure([text.replace(" ", "  ")] * 80)
    assert plain["scope"] == spaced["scope"]
    assert plain["part"] == spaced["part"]
    assert plain["typo"] == spaced["typo"] == 0
    assert tech.measure(["The operator recorded teh reading before checking the instrument."])["typo"] > 0


def test_late_human_typing_traces_are_not_lost_at_the_display_limit():
    r = b.analyze_uploaded_bytes("long.docx", _docx([HUMAN_SLIPPY] * 420))
    assert r.text_source_error_counts["repeated_word"] == 420
    assert r.text_error_shifts[419] < 0
    assert r.score < 52
