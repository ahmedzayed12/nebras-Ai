"""Whole-sentence localisation and geometry; independent of user documents."""
from types import SimpleNamespace as NS

import pytest
import backend_adapter as ba
import nebras_report as nr
import ai_signals
from tests.test_detection_layers import _docx, LLM_PARA


def _report(text, findings=(), features=None, score=95):
    return NS(para_scores=[(0, score, 100)], findings=list(findings),
              features=features or {"ai_lexicon": (1, 1, 1)}, style_spans={})


def _finding(text, marker, kind="ai_phrase"):
    a = text.index(marker)
    return NS(para=0, kind=kind, start=a, end=a + len(marker), note="")


def test_high_paragraph_score_cannot_shade_an_unsupported_sentence_or_gap():
    text = ("This framework provides comprehensive evidence for adaptive interventions across institutions and communities throughout the region. "
            "We measured 18 millimetres on Tuesday and replaced the faulty meter before checking each field reading against our notebook. "
            "The framework provides comprehensive evidence for sustainable interventions across institutions and communities throughout the region.")
    spans = ba.get_engine().split_sentences(text)
    findings = [_finding(text, "This framework"), _finding(text, "The framework")]
    p = nr.plan_highlights(_report(text, findings), [NS(text=text, role="body")])
    assert p.flags[0] == [spans[0], spans[2]]
    assert not p.soft_flags


def test_rsid_and_chat_formatting_cannot_create_a_location():
    import provenance
    text = "We recorded 18 millimetres on Tuesday and replaced the faulty meter before checking the readings again."
    prov = {0: provenance.ParaProv(direct=1.6, rsid="00AA0001")}
    p = nr.plan_highlights(_report(text), [NS(text=text, role="body")], prov)
    assert not p.flags and not p.soft_flags


def test_zeroed_or_cosmetic_evidence_does_not_shade_words():
    text = "First, we inspected the equipment — it needed a battery — before recording measurements on Tuesday."
    findings = [_finding(text, "First", "ai_enum"), _finding(text, "—", "ai_emdash")]
    p = nr.plan_highlights(_report(text, findings), [NS(text=text, role="body")])
    assert not p.flags
    p = nr.plan_highlights(_report(text, [_finding(text, "equipment")],
                                  {"ai_lexicon": (0, 1, 0)}), [NS(text=text, role="body")])
    assert not p.flags


def test_every_supported_sentence_forms_a_complete_paragraph():
    text = ("This framework supports comprehensive interventions across institutions and communities through a coordinated approach to policy development and long term implementation. "
            "These findings underscore sustainable adaptive outcomes across institutions and communities through a coordinated approach to policy development and long term implementation.")
    findings = [_finding(text, "This framework"), _finding(text, "These findings")]
    p = nr.plan_highlights(_report(text, findings), [NS(text=text, role="body")])
    assert p.flags[0] == [(0, len(text))]


def test_filtered_technical_lists_do_not_survive_in_local_scores():
    text = ("Devices use MQTT, CoAP, and Zigbee to transmit the readings from the field to the station. "
            "The network was tested on ARM, MIPS, and RISC-V processors during the trial.")
    r = ba.analyze_uploaded_bytes("protocols.docx", _docx([text] * 10))
    assert not any(f.kind == "ai_tricolon" for f in r.findings)
    assert r.wording_map == ba.get_engine().paragraph_scores(r.doc_paras, [
        f for f in r.findings if f.kind not in {"ai_enum", "ai_emdash", "ai_dash", "ai_markdown"}])
    assert r.score < 30


def test_topic_words_are_removed_from_engine_and_local_vocabulary():
    text = "The framework was revised after twelve trials. Each framework compares the field readings against the laboratory records."
    r = ba.analyze_uploaded_bytes("topic.docx", _docx(["The Framework", *([text] * 10)]))
    assert not any(f.kind == "ai_vocab" and r.texts[f.para][f.start:f.end].lower() == "framework"
                   for f in r.findings)


def test_local_pattern_review_requires_wording_convergence():
    signals = [NS(key="typography", z=1), NS(key="burstiness", z=1), NS(key="repetition", z=1)]
    assert not ba._corroborated_style(NS(sufficient=True, score=95, signals=signals))
    signals += [NS(key="stock", z=1), NS(key="frames", z=1)]
    assert ba._corroborated_style(NS(sufficient=True, score=75, signals=signals))


def test_existing_local_algorithm_can_corroborate_a_missed_pattern():
    eng = ba.get_engine()
    paras = [eng.Para(text=LLM_PARA) for _ in range(12)]
    for p in paras: p.role = "body"
    r = NS(score=20.0, para_scores=[(i, 20.0, 75) for i in range(12)], texts=[p.text for p in paras])
    ba._local_style_review(r, paras)
    assert r.score > 20.0
    assert all(sc > 20 for _, sc, _ in r.para_scores)


def test_arabic_human_account_does_not_open_local_pattern_gate():
    text = "زرنا المحطة يوم الخميس بعد أن تعطلت المضخة. سجلنا ارتفاع الماء بمقدار ١٢ سنتيمترا، ثم اتصلنا بالفني. لم نكمل القراءة الثانية لأن جهاز القياس كان يحتاج إلى بطارية جديدة. أحضرت البطارية صباح الأحد، وأعدنا التجربة مرتين قبل إرسال الجدول إلى المشرف."
    r = ba.analyze_pasted("\n".join([text] * 5))
    assert r.score < 30
    assert not ba._corroborated_style(r.local_style)


def test_local_rhetorical_layer_ignores_named_technical_lists():
    l = ai_signals.analyze(["Devices use MQTT, CoAP, and Zigbee. Neurons have 128, 64, and 32 units."] * 6)
    frames = next(s for s in l.signals if s.key == "frames")
    assert "ثلاثيات: 0" in frames.value


def test_rtl_highlight_rect_covers_the_complete_line(monkeypatch):
    words = [(0, 160, 10, 190, 20, (0, 0, 0)), (0, 120, 10, 150, 20, (0, 0, 0)),
             (0, 80, 10, 110, 20, (0, 0, 0))]
    monkeypatch.setattr(nr, "_pdf_word_stream", lambda *args: (words, {(0, 0, 0): (10, 20)}, [("ا", 0), ("ب", 1), ("ج", 2)]))
    rectangles = []
    dest = NS(draw_rect=lambda rect, **kw: rectangles.append(rect))
    plan = nr.Plan(flags={0: [(0, 5)]})
    matched, total = nr._shade_pdf_pages(None, lambda n: dest, {0: (1, 0, 0)}, ["ا ب ج"], plan)
    assert matched == total == 3
    assert len(rectangles) == 1
    assert rectangles[0].x0 < 80 and rectangles[0].x1 > 190


def test_incomplete_pdf_sentence_is_not_shaded_as_scattered_words(monkeypatch):
    words = [(0, 10, 10, 30, 20, (0, 0, 0))]
    monkeypatch.setattr(nr, "_pdf_word_stream", lambda *args: (words, {(0, 0, 0): (10, 20)}, [("middle", 0)]))
    rectangles = []
    dest = NS(draw_rect=lambda rect, **kw: rectangles.append(rect))
    matched, total = nr._shade_pdf_pages(None, lambda n: dest, {0: (1, 0, 0)},
                                        ["missing middle end"], nr.Plan(flags={0: [(0, 18)]}))
    assert matched == 0 and total == 3
    assert not rectangles


def test_highlight_budget_never_adds_half_a_unit_or_exceeds_its_limit():
    text = " ".join(["reading"] * 100)
    end = len(" ".join(["reading"] * 40))
    plan = nr.Plan(qualifying_words=100, per_para={0: (100, 0)},
                   units=[("strong", 95, 0, 0, end)])
    nr.balance_to_engine_index(NS(score=20), [NS(text=text)], plan)
    assert plan.target_words == 20
    assert plan.total_words <= 30
    assert not plan.flags   # the 40-word unit is omitted whole, never truncated


def test_subject_lists_are_neutral_but_generic_rhetorical_triads_remain():
    import academic_signals as a
    topic = a.topic_terms(["Soil moisture and irrigation control"])
    factual = "The controller measures soil moisture, irrigation pressure, and water flow."
    rhetorical = "The initiative fosters collaboration, resilience, and innovation across sectors."
    assert a.count_triads(factual, topic) == 0
    assert a.count_triads(rhetorical, topic) >= 1
    r = ba.analyze_uploaded_bytes("field.docx", _docx(["Soil moisture and irrigation control", *([factual] * 20)]))
    assert not any(f.kind == "ai_tricolon" for f in r.findings)


def test_proofread_impersonal_measurements_with_tables_are_not_flagged():
    import io
    from docx import Document
    d = Document()
    d.add_heading("Soil moisture measurements", 1)
    texts = [
        "Soil moisture was recorded at 12 stations on 14 April 2024. Each reading was checked against the laboratory record before the pump was started.",
        "The meter at station 7 failed after 18 minutes. Its replacement had a different offset, so the next six observations were excluded from the calculation.",
        "Mean water use was 18 litres per plot in the first trial and 21 litres in the second. The observed difference was smaller than the instrument resolution.",
        "The sampling interval was changed from 15 minutes to 30 minutes after the battery test. Measurements taken before this change were kept in a separate table.",
    ]
    for text in texts * 4: d.add_paragraph(text)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Station"; table.cell(0, 1).text = "Reading"
    table.cell(1, 0).text = "7"; table.cell(1, 1).text = "18 litres"
    data = io.BytesIO(); d.save(data)
    r = ba.analyze_uploaded_bytes("measurements.docx", data.getvalue())
    assert r.score < 30
    p = nr.compute_plan(r, filename="measurements.docx", data=data.getvalue())
    assert not p.flags and not p.soft_flags
